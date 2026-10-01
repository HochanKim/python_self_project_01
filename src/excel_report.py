# ===================================================
# ★ 분석 데이터들 엑셀 내보내기
# ===================================================

import os
import pandas as pd

# 분석 결과 불러오기
from analysis import (
    fault_summary,
    pxl_a_decrib,
    sol_a_decrib,
    luminosity_describe,
    thickness_fault,
    thickness_fault_ratio,
    corr,
    shape_columns_avg,
)

import openpyxl  # 엑셀(.xlsx) 파일을 읽고 수정, 저장을 도와주는 라이브러리
from openpyxl import load_workbook  # openpyxl의 모듈, load_workbook
from openpyxl.utils import get_column_letter  # openpyxl의 모듈, get_column_letter
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
)  # openpyxl의 스타일 기능들 불러오기

# 현재 Python 파일의 위치를 기준으로 reports 폴더 경로 생성
REPORTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "reports")

# Excel 파일 경로
REPORT_FILE = os.path.join(REPORTS, "Steel_Plate_Fault_Analysis.xlsx")

# 헤더 행 번호 저장을 위한 빈 딕셔너리
header_rows = {}


# 헤더 개수가 복수일 경우 사용할 함수 선언
def next_startrow(startrow, dataframe):
    header_size = 1
    return startrow + header_size + len(dataframe) * 2


# ExcelWriter를 이용하여 분석 결과를 각각의 Sheet에 저장
with pd.ExcelWriter(REPORT_FILE, engine="openpyxl") as w:
    # 최종 결함별 결과 보고서
    # fault_summary.to_excel(w, sheet_name="01_Summary", index=False)

    # 결함별 발생 건수 및 비율
    fault_summary.to_excel(w, sheet_name="02_Fault_Count")
    # 헤더가 위치한 행 번호를 딕셔너리에 저장
    header_rows["02_Fault_Count"] = [1]

    # 결함별 Pixels_Areas 기술통계
    pxl_a_decrib.to_excel(w, sheet_name="03_Pixels_Areas")
    # 헤더가 위치한 행 번호를 딕셔너리에 저장
    header_rows["03_Pixels_Areas"] = [1]

    # 결함별 밝기 분석
    ## startrow: 엑셀 저장 위치 지정 (startrow=0 -> 0번째 행부터 데이터 넣기)
    startrow_1 = 0
    startrow_2 = next_startrow(startrow_1, sol_a_decrib)

    sol_a_decrib.to_excel(w, sheet_name="04_Luminosity", startrow=startrow_1)
    luminosity_describe.to_excel(w, sheet_name="04_Luminosity", startrow=startrow_2)
    # 헤더가 위치한 행 번호를 딕셔너리에 저장 (헤더가 복수의 개수가 존재)
    header_rows["04_Luminosity"] = [startrow_1 + 1, startrow_2 + 1]

    # 두께 그룹별 결함 발생 건수
    startrow_1 = 0
    startrow_2 = next_startrow(startrow_1, thickness_fault)

    thickness_fault.to_excel(w, sheet_name="05_Thickness_Groups", startrow=startrow_1)
    thickness_fault_ratio.to_excel(
        w, sheet_name="05_Thickness_Groups", startrow=startrow_2
    )
    # 헤더가 위치한 행 번호를 딕셔너리에 저장 (헤더가 복수의 개수가 존재)
    header_rows["05_Thickness_Groups"] = [startrow_1 + 1, startrow_2 + 1]

    # 상관관계
    corr.to_excel(w, sheet_name="06_Correlation")
    # 헤더가 위치한 행 번호를 딕셔너리에 저장
    header_rows["06_Correlation"] = [1]

    # 형태/방향 변수
    shape_columns_avg.to_excel(w, sheet_name="07_Shape")
    # 헤더가 위치한 행 번호를 딕셔너리에 저장
    header_rows["07_Shape"] = [1]

# 엑셀 통합 문서 불러오기
wb = load_workbook(REPORT_FILE)


for sheet in wb.sheetnames:
    # 시트 이름 지정
    ws = wb[sheet]

    # 숫자 표시 형식
    ## ws.iter_rows(): 행/셀 순회
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.00"

            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"

    # 열 너비 자동 조정
    for column in ws.columns:
        # 현재 열에서 가장 긴 셀 값의 문자 길이를 저장
        max_length = 0

        # 현재 열의 번호를 Excel 열 문자(A, B, C...)로 변환
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            if cell.value is not None:
                # 셀에 값이 있으면 문자열로 변환하여 문자 길이 계산
                cell_length = len(str(cell.value))

                # 지금까지 확인한 값보다 더 긴 값이라면
                if cell_length > max_length:
                    # 가장 긴 문자 길이 갱신
                    max_length = cell_length

        # 가장 긴 문자 길이에 여유 공간 2를 추가하여 열 너비 설정
        ws.column_dimensions[column_letter].width = max_length + 2


# 헤더 꾸미기용 반복문 선언
for sheet_name, rows in header_rows.items():
    ws = wb[sheet_name]

    for row_number in rows:
        for cell in ws[row_number]:
            # 폰트 설정# # plt.show()
            cell.font = Font(bold=True, color="FFFFFF")
            # 셀 배경 설정
            cell.fill = PatternFill(fill_type="solid", fgColor="4472C4")
            # 정렬 설정
            cell.alignment = Alignment(horizontal="center", vertical="center")
            # 테두리 추가
            thin_side = Side(style="thin", color="D9E2F3")
            # 셀 보더 설정
            cell.border = Border(
                left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
            )


# 최종 저장
wb.save(REPORT_FILE)
